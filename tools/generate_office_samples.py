"""Generate original, shareable Office fixtures. Run with uv --with openpyxl --with python-pptx."""
from pathlib import Path
from PIL import Image, ImageDraw
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from pptx import Presentation
from pptx.util import Inches

root = Path('artifacts/office-samples')
root.mkdir(parents=True, exist_ok=True)
picture = Image.new('RGB', (480, 160), 'white')
ImageDraw.Draw(picture).text((20, 40), 'IMAGE ONLY: Invoice ABC-123, Total 42', fill='black', font_size=22)
picture.save(root / 'ocr-candidate.png')

workbook = Workbook()
sheet = workbook.active
sheet.title = '销售汇总'
sheet.merge_cells('A1:C1')
sheet['A1'] = '季度销售报告'
sheet.append(['项目', '金额', '占比'])
sheet.append(['产品 A', 1200, 0.25])
sheet['B3'].number_format = '#,##0.00'
sheet['C3'].number_format = '0.00%'
sheet['B4'] = '=SUM(B3:B3)'
second = workbook.create_sheet('图片附件')
second['A1'] = '下面的图片是待 OCR 区域'
second.add_image(ExcelImage(root / 'ocr-candidate.png'), 'A3')
workbook.save(root / 'sales-with-image.xlsx')

presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[5])
slide.shapes.title.text = '原生幻灯片内容'
text = slide.shapes.add_textbox(Inches(1), Inches(1.3), Inches(8), Inches(1)).text_frame
text.text = '这一段直接从 PPTX 提取，不需要 OCR。'
table = slide.shapes.add_table(3, 2, Inches(1), Inches(2.5), Inches(7), Inches(2)).table
table.cell(0, 0).merge(table.cell(0, 1))
table.cell(0, 0).text = '合并表头'
table.cell(1, 0).text = '项目'
table.cell(1, 1).text = '金额'
table.cell(2, 0).text = '样例'
table.cell(2, 1).text = '42'
slide = presentation.slides.add_slide(presentation.slide_layouts[5])
slide.shapes.title.text = '图片文字待 OCR'
slide.shapes.add_picture(str(root / 'ocr-candidate.png'), Inches(1), Inches(2), width=Inches(8))
presentation.save(root / 'slides-with-table-and-image.pptx')
print(root.resolve())
